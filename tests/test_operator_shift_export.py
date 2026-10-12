"""Regression checks for per-operator confirmed shift exports."""
import importlib.util
import io
import sqlite3
import unittest
import zipfile
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('operator_shift_export', ROOT / 'shift_planner_app' / 'operator_shift_export.py')
export = importlib.util.module_from_spec(spec)
spec.loader.exec_module(export)


class OperatorShiftExportTests(unittest.TestCase):
    def setUp(self):
        self.con = sqlite3.connect(':memory:')
        self.con.row_factory = sqlite3.Row
        self.con.executescript('''
          CREATE TABLE employees(id INTEGER PRIMARY KEY,name TEXT,active INT);
          CREATE TABLE users(id INTEGER PRIMARY KEY,username TEXT,role TEXT,employee_id INT);
          CREATE TABLE shift_patterns(id INTEGER PRIMARY KEY,name TEXT,start_time TEXT,end_time TEXT,break_minutes INT);
          CREATE TABLE assignments(id INTEGER PRIMARY KEY,employee_id INTEGER,work_date TEXT,pattern_id INT,skill TEXT,status TEXT,note TEXT);
          INSERT INTO employees VALUES(1,'山田/太郎',1),(2,'佐藤',0),(3,'未確定者',1);
          INSERT INTO users VALUES(1,'operator_a','operator',1),(2,'operator_b','operator',2);
          INSERT INTO shift_patterns VALUES(1,'早番','09:20','18:00',60),(2,'遅番','14:00','23:00',60);
          INSERT INTO assignments VALUES(1,1,'2026-11-03',1,'consumer','finalized','通常勤務');
          INSERT INTO assignments VALUES(2,1,'2026-11-04',1,'PP','draft','未確定');
          INSERT INTO assignments VALUES(3,1,'2026-12-02',1,'consumer','finalized','別の月');
          INSERT INTO assignments VALUES(4,2,'2026-11-05',2,'merchant','finalized','=CMD()');
          INSERT INTO assignments VALUES(5,3,'2026-11-06',1,'consumer','draft','未確定');
        ''')
        self.con.commit()

    def tearDown(self):
        self.con.close()

    def test_month_and_employee_filter_with_inactive_employee(self):
        result = export.collect_finalized(self.con, '2026-11')
        self.assertEqual(list(result), [2,1])
        self.assertEqual(len(result[1]['assignments']),1)
        self.assertEqual(len(export.collect_finalized(self.con, '2026-11', 2)[2]['assignments']),1)
        self.assertEqual(len(export.collect_finalized(self.con, '2026-11', 3)),0)
        self.assertEqual(len(export.collect_finalized(self.con, '2026-12')),1)

    def test_calendar_disambiguates_unconfirmed_and_off_days(self):
        e = export.collect_finalized(self.con, '2026-11',1)[1]
        rows = export.rows_for_employee(e,'2026-11')
        self.assertEqual(len(rows),30)
        self.assertEqual(rows[2][0],'2026-11-03')
        self.assertEqual(rows[2][4],'確定済み')
        self.assertEqual(rows[2][9],'8時間40分')
        self.assertEqual(rows[3][4],'確定勤務なし')
        self.assertTrue(all('未確定' not in ','.join(map(str, r)) for r in rows))

    def test_xlsx_layout_and_values(self):
        e = export.collect_finalized(self.con,'2026-11',1)[1]
        content = export.build_xlsx(e,'2026-11')
        wb = load_workbook(io.BytesIO(content), read_only=False)
        ws = wb.active
        self.assertEqual(ws.title,'個人別確定シフト')
        self.assertEqual(ws['A1'].value,'2026-11  個人別確定シフト')
        self.assertEqual(ws['A5'].value,'勤務日')
        self.assertEqual(ws['E8'].value,'確定済み')
        self.assertEqual(ws['K8'].value,'consumer')
        self.assertEqual(ws['I8'].value,60)
        self.assertEqual(ws['E9'].value,'確定勤務なし')
        self.assertEqual(ws.max_row,35)
        self.assertEqual(ws.freeze_panes,'C6')
        wb.close()

    def test_csv_bom_and_injection(self):
        e = export.collect_finalized(self.con,'2026-11',2)[2]
        content = export.build_csv(e,'2026-11')
        self.assertTrue(content.startswith(b'\xef\xbb\xbf'))
        text = content.decode('utf-8-sig')
        self.assertIn("'=CMD()", text)
        self.assertNotIn('2026-12-02',text)
        self.assertIn('勤務日,曜日,',text)
        xlsx = export.build_xlsx(e,'2026-11')
        wb=load_workbook(io.BytesIO(xlsx))
        self.assertEqual(wb.active['L10'].data_type,'s')
        self.assertEqual(wb.active['L10'].value,"'=CMD()")
        wb.close()

    def test_zip_separate_files_and_names(self):
        group=export.collect_finalized(self.con,'2026-11')
        zipcontent=export.build_zip(group,'2026-11','xlsx')
        with zipfile.ZipFile(io.BytesIO(zipcontent)) as z:
            names=z.namelist()
            self.assertEqual(len(names),2)
            self.assertEqual(len(set(names)),2)
            self.assertIn('confirmed_shift_2026-11_1_山田_太郎.xlsx',names)
            self.assertIn('confirmed_shift_2026-11_2_佐藤.xlsx',names)
            for name in names:
                self.assertEqual(z.testzip(),None)
                wb=load_workbook(io.BytesIO(z.read(name)))
                self.assertEqual(wb.active['A5'].value,'勤務日')
                wb.close()
        csvzip = export.build_zip(group,'2026-11','csv')
        with zipfile.ZipFile(io.BytesIO(csvzip)) as z:
            self.assertEqual(len(z.namelist()),2)
            self.assertTrue(z.read(z.namelist()[0]).startswith(b'\xef\xbb\xbf'))

    def test_month_validation(self):
        for bad in ('2026-1', '2026-13', '2026-00', '2026-11 and 1=1', '2026-11-01', '', '0-01'):
            with self.assertRaises(ValueError): export.collect_finalized(self.con,bad)
        self.assertEqual(export.safe_filename('\\/|:*?<>"'), '_________')

    def test_db_unchanged(self):
        before=self.con.execute('SELECT id,employee_id,work_date,status FROM assignments ORDER BY id').fetchall()
        export.build_zip(export.collect_finalized(self.con,'2026-11'),'2026-11','xlsx')
        after=self.con.execute('SELECT id,employee_id,work_date,status FROM assignments ORDER BY id').fetchall()
        self.assertEqual([tuple(row) for row in before],[tuple(row) for row in after])


if __name__ == '__main__':
    unittest.main()
