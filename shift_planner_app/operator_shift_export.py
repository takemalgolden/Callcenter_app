"""Administrator-only finalized-shift export utilities (called by Flask route).

Exports one monthly calendar per employee.  Days without a finalized assignment
are explicitly called '確定勤務なし' to avoid silently representing pending drafts
as approved days off.  No data is written back to SQLite.
"""
import calendar
import csv
import io
import re
import zipfile
from collections import OrderedDict
from datetime import date, datetime

WEEKDAYS = ('月', '火', '水', '木', '金', '土', '日')
HEADERS = ('勤務日', '曜日', 'オペレーター', 'ログインID', '確定状況', '勤務パターン',
           '開始時刻', '終了時刻', '休憩(分)', '拘束時間', '当日の主担当スキル', '備考')
EXCEL_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
CSV_MIME = 'text/csv; charset=utf-8'
ZIP_MIME = 'application/zip'


def validate_month(month):
    if not isinstance(month, str) or not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', month):
        raise ValueError('対象月の形式が正しくありません（YYYY-MM）。')
    year, mon = map(int, month.split('-'))
    if year < 1:
        raise ValueError('対象月が正しくありません。')
    return year, mon


def _s(value):
    """Keep CSV/Excel cells from interpreting operator-controlled strings as formulas."""
    s = '' if value is None else str(value)
    if s.lstrip().startswith(('=', '+', '-', '@', '\t', '\r', '\n')):
        return "'" + s
    return s


def safe_filename(name):
    cleaned = re.sub(r'[\\/:*?"<>|\x00-\x1f]', '_', str(name))
    cleaned = cleaned.strip(' .')[:60].rstrip(' .')
    return cleaned or 'operator'


def _duration(start, end):
    if not start or not end:
        return ''
    try:
        sh, sm = map(int, start.split(':'))
        eh, em = map(int, end.split(':'))
        begin, finish = sh * 60 + sm, eh * 60 + em
        if not (0 <= begin < 1440 and 0 <= finish < 1440):
            return ''
        minutes = (finish - begin) % 1440
        return f'{minutes // 60}時間{minutes % 60:02d}分'
    except (ValueError, AttributeError, TypeError):
        return ''


def collect_finalized(conn, month, employee_id=None):
    """Get only employees with finalized shifts in the selected month.

    Preserves deactivated employees and rows with a deleted/null pattern, which
    are still relevant to historical finalization; does not include drafts.
    """
    validate_month(month)
    clause = ' AND e.id=?' if employee_id is not None else ''
    args = (month, employee_id) if employee_id is not None else (month,)
    rows = conn.execute('''
        SELECT a.employee_id, e.name, COALESCE(u.username, '') AS login_id,
               a.work_date, a.skill, a.note, p.name AS pattern_name,
               p.start_time, p.end_time, p.break_minutes
        FROM assignments a
        JOIN employees e ON e.id=a.employee_id
        LEFT JOIN shift_patterns p ON p.id=a.pattern_id
        LEFT JOIN users u ON u.id=(
            SELECT u2.id FROM users u2 WHERE u2.employee_id=e.id
              AND u2.role='operator' ORDER BY u2.id LIMIT 1)
        WHERE substr(a.work_date,1,7)=? AND a.status='finalized' ''' + clause + '''
        ORDER BY e.name, a.employee_id, a.work_date, a.id
    ''', args).fetchall()
    result = OrderedDict()
    for r in rows:
        record = dict(r)
        eid = record['employee_id']
        if eid not in result:
            result[eid] = {'id': eid, 'name': record['name'],
                           'login_id': record['login_id'], 'assignments': []}
        result[eid]['assignments'].append(record)
    return result


def rows_for_employee(employee, month):
    year, mon = validate_month(month)
    assigned = {a['work_date']: a for a in employee['assignments']}
    data = []
    for day in range(1, calendar.monthrange(year, mon)[1] + 1):
        day_value = date(year, mon, day)
        iso = day_value.isoformat()
        shift = assigned.get(iso)
        if shift is None:
            data.append([iso, WEEKDAYS[day_value.weekday()], _s(employee['name']),
                         _s(employee['login_id']), '確定勤務なし', '', '', '', '', '', '', ''])
        else:
            data.append([iso, WEEKDAYS[day_value.weekday()], _s(employee['name']),
                         _s(employee['login_id']), '確定済み',
                         _s(shift['pattern_name']), _s(shift['start_time']),
                         _s(shift['end_time']),
                         shift['break_minutes'] if shift['break_minutes'] is not None else '',
                         _duration(shift['start_time'], shift['end_time']),
                         _s(shift['skill']), _s(shift['note'])])
    return data


def build_csv(employee, month):
    content = io.StringIO(newline='')
    writer = csv.writer(content, lineterminator='\r\n')
    writer.writerow(HEADERS)
    writer.writerows(rows_for_employee(employee, month))
    return ('\ufeff' + content.getvalue()).encode('utf-8')


def build_xlsx(employee, month):
    # Existing application already depends on openpyxl for its Excel exports.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    wb = Workbook()
    ws = wb.active
    ws.title = '個人別確定シフト'
    ws.merge_cells('A1:L1')
    ws['A1'] = f'{month}  個人別確定シフト'
    ws['A1'].font = Font(name='Yu Gothic', size=17, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor='17365D')
    ws['A1'].alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 34
    ws.merge_cells('A2:L2')
    ws['A2'] = f"オペレーター：{_s(employee['name'])}　ログインID：{_s(employee['login_id']) or '未設定'}"
    ws['A2'].font = Font(name='Yu Gothic', size=11, bold=True, color='17365D')
    ws.row_dimensions[2].height = 25
    ws.merge_cells('A3:L3')
    ws['A3'] = f"確定勤務：{len(employee['assignments'])}日　※『確定勤務なし』には、未確定・未登録の日付も含まれます。"
    ws['A3'].font = Font(name='Yu Gothic', size=10, color='536579')
    ws.row_dimensions[3].height = 24
    header_row = 5
    for col, title in enumerate(HEADERS, 1):
        cell = ws.cell(header_row, col, title)
        cell.fill = PatternFill('solid', fgColor='244A74')
        cell.font = Font(name='Yu Gothic', bold=True, color='FFFFFF', size=10)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[header_row].height = 27
    light_border = Side(style='hair', color='DAE2EA')
    data = rows_for_employee(employee, month)
    for offset, values in enumerate(data, header_row + 1):
        for col, value in enumerate(values, 1):
            if col == 1:
                value = date.fromisoformat(value)
            cell = ws.cell(offset, col, value)
            cell.font = Font(name='Yu Gothic', size=10, color='263648')
            cell.alignment = Alignment(vertical='center', wrap_text=col in (6, 11, 12),
                                       horizontal='center' if col in (1, 2, 5, 7, 8, 9, 10) else 'left')
            cell.border = Border(bottom=light_border)
            if col == 1:
                cell.number_format = 'yyyy/mm/dd'
            if values[4] != '確定済み':
                cell.fill = PatternFill('solid', fgColor='F2F5F8')
            elif offset % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F4F9FD')
        ws.row_dimensions[offset].height = 23
    for idx, width in enumerate((15, 7, 19, 19, 14, 17, 13, 13, 12, 13, 20, 30), 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = 'C6'
    ws.auto_filter.ref = f'A{header_row}:L{header_row + len(data)}'
    ws.sheet_view.showGridLines = False
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:5'
    ws.page_margins = PageMargins(left=.25, right=.25, top=.4, bottom=.4, header=.2, footer=.2)
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def export_filename(employee, month, fmt):
    return f"confirmed_shift_{month}_{employee['id']}_{safe_filename(employee['name'])}.{fmt}"


def build_zip(grouped, month, fmt):
    """One Excel/CSV file per employee, no shared sheets or overlapping names."""
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for employee in grouped.values():
            data = build_xlsx(employee, month) if fmt == 'xlsx' else build_csv(employee, month)
            archive.writestr(export_filename(employee, month, fmt), data)
    return output.getvalue()
